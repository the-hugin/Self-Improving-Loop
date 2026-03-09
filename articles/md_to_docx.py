"""
Markdown → DOCX converter using python-docx.
Handles: headings (H1-H3), paragraphs, bold, italic, inline code,
         fenced code blocks, bullet lists, numbered lists, tables, HR.
"""
import sys
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding='utf-8')

INPUT  = r"C:\Users\Petr\.claude\articles\sil-architecture.md"
OUTPUT = r"C:\Users\Petr\.claude\articles\sil-architecture.docx"

# ── Helpers ────────────────────────────────────────────────────────────────

def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'AAAAAA')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(6)


def set_code_block(paragraph):
    paragraph.style = 'No Spacing'
    for run in paragraph.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    # grey background via shading
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F0F0F0')
    pPr.append(shd)


def add_inline_formatted(paragraph, text):
    """Parse **bold**, *italic*, `code` in a line and add runs."""
    # Split on formatting tokens
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    parts = pattern.split(text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            if part:
                paragraph.add_run(part)


def parse_table(doc, table_lines):
    """Parse GFM table lines into a docx table."""
    # Filter separator row (|---|---|)
    rows = [l for l in table_lines if not re.match(r'^\|[-| :]+\|$', l.strip())]
    if not rows:
        return
    parsed = []
    for row in rows:
        cells = [c.strip() for c in row.strip().strip('|').split('|')]
        parsed.append(cells)
    if not parsed:
        return
    ncols = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=ncols)
    table.style = 'Table Grid'
    for i, row in enumerate(parsed):
        for j, cell_text in enumerate(row):
            if j >= ncols:
                break
            cell = table.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            add_inline_formatted(p, cell_text)
            if i == 0:
                for run in p.runs:
                    run.bold = True
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
    doc.add_paragraph()  # spacing after table


# ── Main converter ─────────────────────────────────────────────────────────

def convert(input_path, output_path):
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.2)
        section.right_margin  = Inches(1.2)

    # Base font
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Heading styles
    for i, size in [(1, 20), (2, 16), (3, 13)]:
        h = doc.styles[f'Heading {i}']
        h.font.name = 'Calibri'
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0x1F, 0x35, 0x64) if i == 1 else \
                           RGBColor(0x2E, 0x4D, 0x8A) if i == 2 else \
                           RGBColor(0x37, 0x64, 0x8F)
        h.paragraph_format.space_before = Pt(14 if i == 1 else 10 if i == 2 else 8)
        h.paragraph_format.space_after  = Pt(6)

    with open(input_path, encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    in_code_block = False
    code_lines = []
    table_lines = []
    in_table = False

    def flush_table():
        nonlocal table_lines, in_table
        if table_lines:
            parse_table(doc, table_lines)
        table_lines = []
        in_table = False

    while i < len(lines):
        raw = lines[i].rstrip('\n')
        stripped = raw.strip()

        # ── Fenced code block ──────────────────────────────────────────
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
                i += 1
                continue
            else:
                # End of code block
                in_code_block = False
                if in_table:
                    flush_table()
                full_code = '\n'.join(code_lines)
                for line in code_lines:
                    p = doc.add_paragraph(line)
                    set_code_block(p)
                # thin spacing after block
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
                i += 1
                continue

        if in_code_block:
            code_lines.append(raw)
            i += 1
            continue

        # ── Table ─────────────────────────────────────────────────────
        if stripped.startswith('|'):
            if in_table:
                table_lines.append(stripped)
            else:
                if in_table:
                    flush_table()
                in_table = True
                table_lines = [stripped]
            i += 1
            continue
        else:
            if in_table:
                flush_table()

        # ── Blank line ────────────────────────────────────────────────
        if not stripped:
            i += 1
            continue

        # ── Horizontal rule ──────────────────────────────────────────
        if re.match(r'^---+$', stripped) or re.match(r'^\*\*\*+$', stripped):
            add_horizontal_rule(doc)
            i += 1
            continue

        # ── Headings ─────────────────────────────────────────────────
        m = re.match(r'^(#{1,3})\s+(.*)', stripped)
        if m:
            level = len(m.group(1))
            text  = m.group(2)
            # Strip anchor links like {#anchor}
            text = re.sub(r'\{#[^}]+\}', '', text).strip()
            p = doc.add_heading(level=level)
            p.clear()
            add_inline_formatted(p, text)
            i += 1
            continue

        # ── Bullet list ───────────────────────────────────────────────
        m = re.match(r'^(\s*)[-*]\s+(.*)', raw)
        if m:
            indent = len(m.group(1)) // 2
            text   = m.group(2)
            style_name = 'List Bullet 2' if indent > 0 else 'List Bullet'
            p = doc.add_paragraph(style=style_name)
            add_inline_formatted(p, text)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # ── Numbered list ─────────────────────────────────────────────
        m = re.match(r'^(\s*)\d+\.\s+(.*)', raw)
        if m:
            text = m.group(2)
            p = doc.add_paragraph(style='List Number')
            add_inline_formatted(p, text)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # ── Blockquote ────────────────────────────────────────────────
        if stripped.startswith('>'):
            text = stripped.lstrip('> ').strip()
            p = doc.add_paragraph(style='Quote') if 'Quote' in doc.styles \
                else doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            for run in p.runs:
                pass
            add_inline_formatted(p, text)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # ── Normal paragraph ──────────────────────────────────────────
        p = doc.add_paragraph()
        add_inline_formatted(p, stripped)
        p.paragraph_format.space_after = Pt(6)
        i += 1

    if in_table:
        flush_table()

    doc.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == '__main__':
    convert(INPUT, OUTPUT)
